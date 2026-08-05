#!/usr/bin/env python3
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'outputs'; FIG=OUT/'figures'; m=json.loads((OUT/'metrics.json').read_text()); inv=json.loads((OUT/'inverse_design.json').read_text())
NAVY='082B4C'; TEAL='008C95'; GOLD='D6A33A'; PALE='EAF4F4'; GRAY='536777'; WHITE='FFFFFF'
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.82); sec.right_margin=Inches(.82); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
for name,size,color,bold,bef,aft in [('Normal',10.2,'263746',False,0,6),('Title',28,NAVY,True,0,8),('Subtitle',13,GRAY,False,0,12),('Heading 1',17,NAVY,True,14,6),('Heading 2',13,TEAL,True,10,4),('Heading 3',11,GOLD,True,8,3)]:
 s=doc.styles[name]; s.font.name='Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=bold; s.paragraph_format.space_before=Pt(bef); s.paragraph_format.space_after=Pt(aft); s.paragraph_format.line_spacing=1.12
h=sec.header.paragraphs[0]; h.text='PROCESSTWIN  |  TECHNICAL REPORT'; h.runs[0].font.size=Pt(8); h.runs[0].font.bold=True; h.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; f.add_run('Harsh Saand  |  Physics-generated prototype evidence  |  '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
for r in f.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
def shade(c,fill):
 p=c._tc.get_or_add_tcPr(); x=p.find(qn('w:shd')) or OxmlElement('w:shd'); x.set(qn('w:fill'),fill); p.append(x) if x.getparent() is None else None
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for i,x in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=str(x); shade(c,NAVY)
  for z in c.paragraphs[0].runs: z.font.bold=True; z.font.color.rgb=RGBColor.from_string(WHITE); z.font.size=Pt(8.5)
 for j,row in enumerate(rows):
  cs=t.add_row().cells
  for i,x in enumerate(row):
   cs[i].text=str(x); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if j%2: shade(cs[i],'F3F7F9')
   for z in cs[i].paragraphs[0].runs: z.font.size=Pt(8.3)
 for row in t.rows:
  for c,w in zip(row.cells,widths): c.width=Inches(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
def pic(name,width,caption):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(FIG/name),width=Inches(width)); q=doc.add_paragraph(caption); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.runs[0].italic=True; q.runs[0].font.size=Pt(8.3); q.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
def callout(label,text):
 t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,PALE); p=c.paragraphs[0]; a=p.add_run(label+'  '); a.bold=True; a.font.color.rgb=RGBColor.from_string(TEAL); p.add_run(text); doc.add_paragraph().paragraph_format.space_after=Pt(1)
def bullet(x): doc.add_paragraph(x,style='List Bullet')

doc.add_paragraph('A*STAR IME PORTFOLIO PROTOTYPE').runs[0].font.color.rgb=RGBColor.from_string(TEAL)
doc.add_paragraph('ProcessTwin',style='Title'); doc.add_paragraph('Physics-grounded AI surrogates for silicon thermal oxidation, dopant diffusion, and inverse recipe design',style='Subtitle'); doc.add_paragraph('HARSH SAAND',style='Heading 2'); doc.add_paragraph('Technical report  |  30 July 2026  |  Locally measured on Apple M3 Pro CPU')
pic('system_flow.png',4.5,'Implemented physics-to-AI workflow.')
callout('Evidence boundary','All quantitative evidence measures surrogate fidelity to the declared reduced-order simulator. This is not commercial process software, fab calibration, or production process prediction.')
doc.add_page_break(); doc.add_heading('Executive summary',1)
t=m['test']; o=m['ood']; doc.add_paragraph(f"ProcessTwin implements a complete AI-for-process-modeling workflow: transparent silicon oxidation and dopant-diffusion physics, a 6,000-recipe design of experiments, blocked recipe-region evaluation, an interpretable polynomial/PCA baseline, a three-member neural/PCA surrogate, uncertainty analysis, full-profile reconstruction, and simulator-verified inverse recipe design. On 477 held-out temperature-bin recipes, the neural surrogate achieved oxide-thickness MAE {t['oxide_nm']['mae']:.2f} nm, junction-depth MAE {t['junction_um']['mae']:.4f} µm, and log-profile RMSE {t['profile_log10_rmse']:.3f} decades.")
callout('Measured contribution',f"Compared with the baseline, the neural surrogate reduced held-out oxide MAE from {t['baseline']['oxide_nm']['mae']:.2f} to {t['oxide_nm']['mae']:.2f} nm, junction MAE from {t['baseline']['junction_um']['mae']:.4f} to {t['junction_um']['mae']:.4f} µm, and profile error from {t['baseline']['profile_log10_rmse']:.3f} to {t['profile_log10_rmse']:.3f} decades.")
doc.add_heading('Semiconductor process problem',1)
doc.add_paragraph('Thermal oxidation controls gate and isolation dielectric thickness, while implant and anneal conditions control dopant-profile shape and junction depth. Process exploration repeatedly evaluates temperature, time, ambient, species, dose, projected range, straggle, and substrate background. A learned surrogate can screen this multidimensional space quickly, but only when its physical assumptions, interpolation domain, and error are explicit.')
pic('system_flow.png',4.15,'Figure 1. Implemented ProcessTwin system flow.')

doc.add_page_break(); doc.add_heading('Reduced-order process physics',1)
doc.add_heading('Deal-Grove oxidation',2); doc.add_paragraph('Oxide thickness x follows x² + Ax = B(t + τ). The implementation exposes dry/wet reference coefficients and applies Arrhenius temperature scaling to B and B/A. Initial oxide is included through the time-shift equivalent. The resulting response is monotonic in time and strongly dependent on temperature and ambient.')
doc.add_heading('Dopant diffusion',2); doc.add_paragraph('A Gaussian implant profile is parameterized by dose, projected range, and straggle. Constant intrinsic diffusivity follows D = D₀ exp(-Ea/kT) for boron or phosphorus. Annealing broadens the variance by 2Dt while conserving integrated dose in the infinite-domain analytical solution. Junction depth is the deepest interpolated crossing of the declared background concentration.')
table(['Physics block','Inputs','Outputs','Explicit omissions'],[
 ('Oxidation','Ambient, T, time, initial oxide','Oxide thickness','Orientation, stress, dopant effects, equipment calibration'),
 ('Implant profile','Species, dose, Rp, straggle','Initial concentration vs depth','Channeling, damage, Pearson-IV tails'),
 ('Anneal diffusion','T, time, background','Final profile, peak, junction, retained dose','TED, clustering, activation, segregation')],[1.2,1.7,1.7,2.0])
pic('oxidation_surface.png',6.0,'Figure 2. Wet-oxidation response surface; the dashed line marks the primary temperature-domain boundary.')

doc.add_page_break(); doc.add_heading('Simulation dataset and evaluation design',1)
doc.add_paragraph('A deterministic seed-42 DOE generated 6,000 recipes and 121-point depth profiles from 0 to 1.2 µm. The primary domain contains oxidation temperatures of 850–1050 °C and anneal temperatures of 800–1050 °C. Entire 25 °C anneal-temperature bins were assigned to validation or test, preventing neighbouring temperature recipes from being randomly scattered across partitions. A separate 960-recipe OOD set uses anneal temperatures of 1055–1150 °C and includes oxidation temperatures up to 1150 °C.')
table(['Partition','Recipes','Purpose'],[('Train','4,105','Model fitting'),('Validation','458','Internal early stopping'),('Held-out test','477','Blocked-bin interpolation evidence'),('High-temperature OOD','960','Extrapolation stress test')],[2.0,1.1,3.4])
doc.add_heading('Recipe ranges',2)
table(['Variable','Range'],[('Oxidation temperature','850–1050 °C primary; up to 1150 °C OOD'),('Oxidation time','5–180 min'),('Ambient','dry / wet'),('Anneal temperature','800–1050 °C primary; 1055–1150 °C OOD'),('Anneal time','5–1000 s, log-uniform'),('Species','B / P'),('Dose','10^12.5–10^15.5 cm^-2'),('Projected range / straggle','0.03–0.35 µm / 0.015–0.09 µm'),('Background','10^15–10^17 cm^-3')],[2.4,4.1])
doc.add_paragraph('The dataset is physics-generated; it contains no experimental wafers, proprietary recipes, or claimed fab measurements. Every result is reproducible from the published simulator configuration.')

doc.add_page_break(); doc.add_heading('AI methodology',1)
doc.add_heading('Baseline',2); doc.add_paragraph('The baseline standardizes recipe features, expands them to second-order polynomial terms, and fits ridge regression. Full log-concentration profiles are compressed into ten principal components and reconstructed after coefficient prediction. This provides a fast, interpretable reference with modest nonlinear capacity.')
doc.add_heading('Neural surrogate',2); doc.add_paragraph('The main model is an ensemble of three independently seeded multilayer perceptrons with 128–128–64 hidden units, ReLU activations, standardized inputs and standardized multi-task outputs. It jointly predicts oxide thickness, junction depth, log peak concentration, and ten PCA profile coefficients. Early stopping uses an internal training-only holdout. Ensemble spread provides a prototype epistemic-uncertainty signal.')
table(['Output','Loss representation','Evaluation'],[('Oxide thickness','nm, standardized target','MAE, RMSE, R²'),('Junction depth','µm, standardized target','MAE, RMSE, R²'),('Peak concentration','log10 cm^-3','MAE, RMSE, R²'),('Full profile','10 PCA coefficients of log10 profile','Depthwise log10 RMSE'),('Uncertainty','Three-member prediction spread','Empirical 90% interval coverage')],[1.35,2.7,2.45])
pic('profile_overlay.png',6.1,'Figure 3. Representative held-out physics and surrogate concentration profiles.')

doc.add_page_break(); doc.add_heading('Locally measured results',1)
table(['Quantity','Baseline','Neural surrogate','Neural R²'],[
 ('Oxide MAE',f"{t['baseline']['oxide_nm']['mae']:.2f} nm",f"{t['oxide_nm']['mae']:.2f} nm",f"{t['oxide_nm']['r2']:.4f}"),
 ('Junction MAE',f"{t['baseline']['junction_um']['mae']:.4f} µm",f"{t['junction_um']['mae']:.4f} µm",f"{t['junction_um']['r2']:.4f}"),
 ('Log-peak MAE',f"{t['baseline']['log10_peak']['mae']:.4f}",f"{t['log10_peak']['mae']:.4f}",f"{t['log10_peak']['r2']:.4f}"),
 ('Profile RMSE',f"{t['baseline']['profile_log10_rmse']:.3f} decades",f"{t['profile_log10_rmse']:.3f} decades",'—')],[1.55,1.7,1.8,1.1])
pic('model_comparison.png',6.5,'Figure 4. Baseline and neural-surrogate errors on held-out recipe bins.')
pic('parity.png',6.4,'Figure 5. Held-out simulator-versus-surrogate parity.')
doc.add_paragraph(f"Batched neural inference measured {t['inference_us_per_recipe']:.2f} µs per recipe, compared with {t['physics_us_per_recipe']:.2f} µs for this already-lightweight analytical simulator ({t['measured_speedup']:.2f}× measured ratio). The runtime result is not presented as a dramatic speedup claim: the process physics is deliberately compact. The value is that the workflow scales to costlier solvers while retaining explicit error and domain checks.")

doc.add_page_break(); doc.add_heading('OOD and uncertainty findings',1)
table(['Metric','Held-out test','High-temperature OOD'],[
 ('Oxide MAE',f"{t['oxide_nm']['mae']:.2f} nm",f"{o['oxide_nm']['mae']:.2f} nm"),('Junction MAE',f"{t['junction_um']['mae']:.4f} µm",f"{o['junction_um']['mae']:.4f} µm"),('Profile RMSE',f"{t['profile_log10_rmse']:.3f}",f"{o['profile_log10_rmse']:.3f}"),('Oxide 90% coverage',f"{t['ensemble_90_coverage']['oxide_nm']:.3f}",f"{o['ensemble_90_coverage']['oxide_nm']:.3f}"),('Junction 90% coverage',f"{t['ensemble_90_coverage']['junction_um']:.3f}",f"{o['ensemble_90_coverage']['junction_um']:.3f}")],[2.4,1.8,2.0])
doc.add_paragraph('Errors rise materially outside the training temperature domain, especially for full profiles. Empirical ensemble intervals also under-cover the nominal 90% target. The interface therefore flags high-temperature inputs as review-only. Ensemble spread is useful triage evidence but is not fully calibrated uncertainty.')

doc.add_page_break(); doc.add_heading('Inverse recipe design',1)
doc.add_paragraph('A constrained differential-evolution search uses the neural surrogate to propose a wet-oxidation/boron-anneal recipe for targets of 120 nm oxide and 0.32 µm junction depth. The candidate is always checked in the physics simulator. Because the initial candidate missed oxide thickness, a bounded local simulator refinement was applied; both stages are retained in the artifact rather than hiding the correction.')
pre=inv['candidate_solver_check']; fin=inv['refined_solver_verified']; rr=inv['refined_recipe']
table(['Stage','Oxide','Junction'],[('Target','120.000 nm','0.320000 µm'),('Surrogate candidate, simulator check',f"{pre['oxide_nm']:.3f} nm",f"{pre['junction_um']:.6f} µm"),('Simulator-refined final',f"{fin['oxide_nm']:.6f} nm",f"{fin['junction_um']:.9f} µm")],[2.8,1.8,1.8])
table(['Final recipe variable','Value'],[('Wet oxidation temperature',f"{rr['oxidation_temp_c']:.2f} °C"),('Wet oxidation time',f"{rr['oxidation_time_min']:.2f} min"),('Boron anneal temperature',f"{rr['anneal_temp_c']:.2f} °C"),('Anneal time',f"{rr['anneal_time_s']:.2f} s"),('Dose / Rp / straggle','1e14 cm^-2 / 0.12 µm / 0.04 µm')],[3.0,3.4])
callout('Interpretation','The miss-and-refine sequence is valuable evidence: surrogate optimization is fast, but simulator verification remains necessary. This is a software validation result, not a manufacturable recipe recommendation.')
doc.add_heading('Engineer-facing interpretability',1)
for x in ['Response surfaces expose monotonic temperature/time behaviour.','Physics-versus-AI profile overlays show where reconstruction error occurs in depth.','Parity plots reveal bias, outliers, and dynamic range for scalar endpoints.','Baseline comparison separates learned improvement from simple polynomial interpolation.','OOD error and ensemble coverage quantify when review is required.']: bullet(x)

doc.add_heading('Limitations and validity threats',1)
for x in ['Simplified 1D continuum physics with literature-scale coefficients; no experimental calibration.','Gaussian implant approximation omits channeling, damage and asymmetric tails.','Constant intrinsic diffusivity omits concentration dependence, TED, clustering, activation and segregation.','Deal-Grove approximation omits orientation, stress, dopant and thin-oxide corrections.','Blocked temperature bins reduce but do not eliminate recipe-space correlation.','Only three ensemble members and under-covered intervals; uncertainty is not fab-grade.','Inverse design is verified only against the same reduced-order simulator.','No claim is made about commercial process-software equivalence or fabrication performance.']: bullet(x)

doc.add_page_break(); doc.add_heading('Conclusion',1)
doc.add_paragraph(f"ProcessTwin demonstrates a complete, locally measured AI-for-process-modeling prototype. On held-out recipe bins, the neural surrogate reached {t['oxide_nm']['mae']:.2f} nm oxide MAE, {t['junction_um']['mae']:.4f} µm junction MAE, and {t['profile_log10_rmse']:.3f}-decade profile RMSE, improving materially on the polynomial/PCA baseline for oxide, junction, and profile reconstruction. The separate high-temperature test showed the expected deterioration and justified explicit OOD routing.")
doc.add_paragraph('The project’s strongest evidence is methodological: declared physics, reproducible simulation data, blocked evaluation, a real baseline, multi-output neural modeling, uncertainty checks, interpretable profile diagnostics, and mandatory simulator verification of inverse-designed recipes. Its limitations are equally explicit, keeping prototype results separate from fab or commercial process-software claims.')
doc.add_heading('References and parameter provenance',1)
for x in ['B. E. Deal and A. S. Grove, “General Relationship for the Thermal Oxidation of Silicon,” Journal of Applied Physics 36, 3770–3778 (1965).','S. M. Sze and K. K. Ng, Physics of Semiconductor Devices, 3rd ed., Wiley (2006).','J. D. Plummer, M. D. Deal, and P. B. Griffin, Silicon VLSI Technology, Prentice Hall (2000).','The diffusivity and oxidation coefficients in src/processtwin/physics.py are exposed literature-scale approximations; they are not equipment-calibrated parameters.']: bullet(x)
out=OUT/'processtwin_report.docx'; doc.save(out); print(out)
